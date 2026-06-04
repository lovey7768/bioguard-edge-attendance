import sqlite3
import os
from datetime import datetime

class AttendanceDatabase:
    def __init__(self, db_path="data/attendance.db"):
        self.db_path = db_path
        # Ensure the directory exists
        os.makedirs(os.path.dirname(self.db_path), exist_ok=True)
        self.init_tables()

    def _get_connection(self):
        """Returns a unique connection instance for thread-safe operations."""
        return sqlite3.connect(self.db_path)

    def init_tables(self):
        """Initializes relational schemas for user profiles and logs."""
        with self._get_connection() as conn:
            cursor = conn.cursor()
            
            # 1. Users Profile Table
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS users (
                    user_id TEXT PRIMARY KEY,
                    name TEXT NOT NULL,
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
            """)
            
            # 2. Daily Attendance Logs Table
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS attendance_logs (
                    log_id INTEGER PRIMARY KEY AUTOINCREMENT,
                    user_id TEXT,
                    date TEXT NOT NULL,
                    timestamp TEXT NOT NULL,
                    check_out_time TEXT DEFAULT NULL,
                    status TEXT DEFAULT 'Present',
                    FOREIGN KEY (user_id) REFERENCES users(user_id),
                    UNIQUE(user_id, date)
                )
            """)
            # Migrate: add check_out_time column if it doesn't exist yet
            try:
                cursor.execute("ALTER TABLE attendance_logs ADD COLUMN check_out_time TEXT DEFAULT NULL")
            except Exception:
                pass  # Column already exists
            conn.commit()

    def register_user(self, user_id, name):
        """Inserts a new authorized face profile record into the system."""
        try:
            with self._get_connection() as conn:
                cursor = conn.cursor()
                cursor.execute(
                    "INSERT INTO users (user_id, name) VALUES (?, ?)", 
                    (str(user_id).strip(), name.strip())
                )
                conn.commit()
                print(f"[OK] Successfully registered user: {name} ({user_id})")
                return True
        except sqlite3.IntegrityError:
            print(f"[WARN] User ID '{user_id}' already exists in the database.")
            return False

    def log_attendance(self, user_id):
        """Logs a timestamped record if the user hasn't checked in today yet."""
        now = datetime.now()
        current_date = now.strftime("%Y-%m-%d")
        current_time = now.strftime("%H:%M:%S")
        
        try:
            with self._get_connection() as conn:
                cursor = conn.cursor()
                
                # Fetch name to return a clean confirmation message
                cursor.execute("SELECT name FROM users WHERE user_id = ?", (user_id,))
                user = cursor.fetchone()
                if not user:
                    print(f"[ERROR] User ID '{user_id}' is not registered.")
                    return None
                
                name = user[0]
                
                # Insert log (will fail if UNIQUE constraint on user_id + date triggers)
                cursor.execute("""
                    INSERT INTO attendance_logs (user_id, date, timestamp) 
                    VALUES (?, ?, ?)
                """, (user_id, current_date, current_time))
                conn.commit()
                
                print(f"[OK] Marked Present: {name} at {current_time}")
                return {"user_id": user_id, "name": name, "time": current_time, "new_log": True}
                
        except sqlite3.IntegrityError:
            # Handles duplicate daily scans gracefully (already checked in today)
            with self._get_connection() as conn2:
                cur2 = conn2.cursor()
                cur2.execute("SELECT name FROM users WHERE user_id = ?", (user_id,))
                user = cur2.fetchone()
            name = user[0] if user else user_id
            return {"user_id": user_id, "name": name, "time": current_time, "new_log": False}

    def log_checkout(self, user_id):
        """Records the check-out time for a user who has already checked in today."""
        now = datetime.now()
        current_date = now.strftime("%Y-%m-%d")
        current_time = now.strftime("%H:%M:%S")

        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT name FROM users WHERE user_id = ?", (user_id,))
            user = cursor.fetchone()
            if not user:
                print(f"[ERROR] User ID '{user_id}' is not registered.")
                return None
            name = user[0]

            # Check if a login record exists for today
            cursor.execute(
                "SELECT log_id FROM attendance_logs WHERE user_id = ? AND date = ?",
                (user_id, current_date)
            )
            record = cursor.fetchone()

            if record:
                cursor.execute(
                    "UPDATE attendance_logs SET check_out_time = ? WHERE log_id = ?",
                    (current_time, record[0])
                )
                conn.commit()
                print(f"[OK] Checked Out: {name} at {current_time}")
                return {"user_id": user_id, "name": name, "time": current_time, "checked_out": True}
            else:
                print(f"[WARN] {name} has no check-in record today — cannot check out.")
                return {"user_id": user_id, "name": name, "checked_out": False}

    def fetch_todays_attendance(self):
        """Retrieves all attendance entries marked for the current calendar date."""
        current_date = datetime.now().strftime("%Y-%m-%d")
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                SELECT u.user_id, u.name, a.timestamp 
                FROM attendance_logs a
                JOIN users u ON a.user_id = u.user_id
                WHERE a.date = ?
                ORDER BY a.timestamp DESC
            """, (current_date,))
            return cursor.fetchall()

    def export_to_excel(self, output_path="data/attendance.xlsx"):
        """Exports all attendance logs to a formatted Excel file."""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import (Font, PatternFill, Alignment,
                                         Border, Side)
            from openpyxl.utils import get_column_letter
        except ImportError:
            print("[ERROR] openpyxl not installed. Run: pip install openpyxl")
            return None

        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                SELECT u.name, u.user_id, a.date, a.timestamp, a.check_out_time, a.status
                FROM attendance_logs a
                JOIN users u ON a.user_id = u.user_id
                ORDER BY a.date DESC, a.timestamp DESC
            """)
            rows = cursor.fetchall()

        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance Log"

        # ── Styles ────────────────────────────────────────────────────────
        header_font    = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
        header_fill    = PatternFill("solid", fgColor="1F4E79")   # dark blue
        alt_fill       = PatternFill("solid", fgColor="D9E1F2")   # light blue
        center_align   = Alignment(horizontal="center", vertical="center")
        thin_border    = Border(
            left   = Side(style="thin", color="BFBFBF"),
            right  = Side(style="thin", color="BFBFBF"),
            top    = Side(style="thin", color="BFBFBF"),
            bottom = Side(style="thin", color="BFBFBF"),
        )

        # ── Title row ─────────────────────────────────────────────────────
        ws.merge_cells("A1:F1")
        title_cell = ws["A1"]
        title_cell.value     = "Smart Attendance System — Full Log"
        title_cell.font      = Font(name="Calibri", bold=True, size=14, color="1F4E79")
        title_cell.alignment = center_align
        ws.row_dimensions[1].height = 28

        # ── Sub-title (export timestamp) ──────────────────────────────────
        ws.merge_cells("A2:F2")
        sub_cell = ws["A2"]
        sub_cell.value     = f"Exported on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        sub_cell.font      = Font(name="Calibri", italic=True, size=10, color="808080")
        sub_cell.alignment = center_align
        ws.row_dimensions[2].height = 18

        # ── Headers ───────────────────────────────────────────────────────
        headers = ["Name", "User ID", "Date", "Check-In Time", "Check-Out Time", "Status"]
        col_widths = [22, 16, 14, 16, 16, 12]

        for col, (header, width) in enumerate(zip(headers, col_widths), start=1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font      = header_font
            cell.fill      = header_fill
            cell.alignment = center_align
            cell.border    = thin_border
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.row_dimensions[3].height = 22

        # ── Data rows ─────────────────────────────────────────────────────
        for row_idx, record in enumerate(rows, start=4):
            fill = alt_fill if row_idx % 2 == 0 else PatternFill()
            for col, value in enumerate(record, start=1):
                cell = ws.cell(row=row_idx, column=col, value=value if value is not None else "-")
                cell.font      = Font(name="Calibri", size=11)
                cell.fill      = fill
                cell.alignment = center_align
                cell.border    = thin_border

        # ── Freeze header rows ────────────────────────────────────────────
        ws.freeze_panes = "A4"

        # ── Summary sheet ─────────────────────────────────────────────────
        ws2 = wb.create_sheet("Summary")
        ws2["A1"] = "Total Records"
        ws2["B1"] = len(rows)
        ws2["A1"].font = Font(bold=True)

        unique_dates = sorted(set(r[2] for r in rows), reverse=True)
        ws2["A3"] = "Date"
        ws2["B3"] = "Present Count"
        ws2["A3"].font = ws2["B3"].font = Font(bold=True, color="FFFFFF")
        ws2["A3"].fill = ws2["B3"].fill = PatternFill("solid", fgColor="1F4E79")
        for i, date in enumerate(unique_dates, start=4):
            count = sum(1 for r in rows if r[2] == date)
            ws2.cell(row=i, column=1, value=date)
            ws2.cell(row=i, column=2, value=count)
        ws2.column_dimensions["A"].width = 14
        ws2.column_dimensions["B"].width = 16

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        try:
            wb.save(output_path)
            print(f"[OK] Attendance exported to: {output_path}")
            return output_path
        except PermissionError:
            # File is open in Excel — save a timestamped backup instead
            stamp  = datetime.now().strftime("%Y%m%d_%H%M%S")
            backup = output_path.replace(".xlsx", f"_{stamp}.xlsx")
            wb.save(backup)
            print(f"[WARN] '{output_path}' is open in Excel. Saved backup to: {backup}")
            return backup